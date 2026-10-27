"""Build a filterable companion index from the reviewed research records.

This is an index, not an automatic adjudication of scientific claims.
It does not change source papers, dossiers, or assignment records.
"""
from __future__ import annotations

import json
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent


def read(name, default=None):
    path = ROOT / name
    return json.loads(path.read_text(encoding="utf-8-sig")) if path.exists() else default


def display(value):
    if value is None:
        return ""
    if isinstance(value, list) and all(isinstance(v, str) for v in value):
        value = "; ".join(value)
    elif isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    if isinstance(value, str):
        if len(value) > 32000:
            value = value[:31900] + " … [truncated in workbook; consult linked JSON/dossier]"
        if value.startswith(("=", "+", "-", "@")):
            value = "'" + value
    return value


def sheet(book, title, rows, fields, widths=None):
    ws = book.create_sheet(title)
    ws.append(fields)
    for row in rows:
        ws.append([display(row.get(field)) for field in fields])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="173C50")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[1].height = 30
    for i, field in enumerate(fields, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(field, 36)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F1F5F7")
        ws.row_dimensions[row[0].row].height = 62
    return ws


def main():
    reviewed = read("inventory/reviewed_papers.json", [])
    ledger = read("assignment_ledger.json", [])
    book = Workbook()
    book.remove(book.active)
    info = [
        {"Field": "Literature cutoff", "Value": "2026-10-22"},
        {"Field": "Generated UTC", "Value": datetime.now(timezone.utc).isoformat()},
        {"Field": "Purpose", "Value": "Filter papers, source claims, bottlenecks and assignment status. Read linked dossiers for interpretation, exact locators and evidence/access limits."},
        {"Field": "Coverage", "Value": f"{len(reviewed)} inventory entries; {sum(x.get('status') == 'Complete' for x in ledger)} of 10 delegated assignments marked complete at generation."},
        {"Field": "Evidence", "Value": "Author reports, inspected code, executed isolated checks and hypotheses are distinct. This index does not establish benchmark reproduction or scientific validity."},
        {"Field": "Folders", "Value": "The supplied Paper folder and its ten grouping dimensions remain separate from this research record. A paper can belong to multiple groups."},
        {"Field": "Definitions", "Value": "VLN-CE means continuous environments; CE loss means cross-entropy. Training-free inference does not imply an untrained model or absence of persistent feedback."},
    ]
    sheet(book, "Read me", info, ["Field", "Value"], {"Field": 25, "Value": 110})
    paper_rows = []
    for p in reviewed:
        row = {k: p.get(k) for k in ("paper_id", "short_name", "full_title", "authors", "year", "venue_or_status", "assignment", "completion_status", "repository_status", "repo_url", "repo_commit", "dossier_path", "access_limitations")}
        row["inspected_code_entries"] = len(p.get("code_files_inspected", []))
        row["reviewer_code_executed"] = p.get("code_executed", False)
        row["pdf_files"] = len(p.get("local_files", []))
        row["execution_scope"] = p.get("code_execution_field_scope", "")
        row["coordinator_executed_components"] = p.get("coordinator_executed_components", [])
        row["independent_evaluator_executed_components"] = p.get("independent_evaluator_executed_components", [])
        paper_rows.append(row)
    paper_fields = ["short_name", "full_title", "authors", "year", "venue_or_status", "assignment", "completion_status", "repository_status", "inspected_code_entries", "reviewer_code_executed", "repo_url", "repo_commit", "pdf_files", "dossier_path", "access_limitations", "execution_scope", "coordinator_executed_components", "independent_evaluator_executed_components", "paper_id"]
    ws = sheet(book, "Papers", paper_rows, paper_fields, {"short_name": 24, "full_title": 58, "authors": 45, "year": 10, "assignment": 12, "inspected_code_entries": 15, "pdf_files": 12, "access_limitations": 65})
    for i, p in enumerate(paper_rows, 2):
        for field in ("dossier_path", "repo_url"):
            value = p.get(field)
            if value:
                cell = ws.cell(i, paper_fields.index(field) + 1)
                cell.hyperlink = value
                cell.font = Font(color="0563C1", underline="single")

    specs = [
        ("Claims", "evidence/claims.json", ["claim_id", "paper_id", "assignment", "claim", "evidence_type", "source", "locator"]),
        ("Paper bottlenecks", "evidence/paper_bottlenecks.json", ["id", "paper_id", "description", "conditions", "cause", "evidence", "kind", "followups", "status", "confidence"]),
        ("Follow-up matrix", "followups/Bottleneck_Matrix.json", None),
        ("Follow-up sources", "followups/Source_Index.json", ["source_id", "title", "authors", "year", "venue_or_status", "primary_url", "provided_paper_id", "reviewed_pages", "repo_url", "repo_commit", "code_files_inspected", "limitations"]),
        ("Proposals", "proposals/final_proposals.json" if (ROOT / "proposals/final_proposals.json").exists() else "proposals/proposals.json", None),
        ("Final ranking", "proposals/final_ranking.json", None),
        ("Research gaps", "proposals/research_gaps.json", None),
        ("Assignments", "assignment_ledger.json", ["assignment", "scope", "status", "agent", "paper_ids", "outputs", "started_utc", "completed_utc"]),
    ]
    for title, source, fields in specs:
        rows = read(source, [])
        if isinstance(rows, dict):
            rows = rows.get("rows", rows.get("proposals", [rows]))
        if not rows:
            continue
        if fields is None:
            fields = list(dict.fromkeys(key for row in rows for key in row))
        sheet(book, title, rows, fields)
    target = ROOT / "Research_Index.xlsx"
    book.save(target)
    check = load_workbook(target, read_only=True, data_only=False)
    summary = {ws.title: {"rows": ws.max_row - 1, "columns": ws.max_column} for ws in check.worksheets}
    check.close()
    (ROOT / "appendices/workbook_validation.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"file": str(target), "sheets": summary}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
